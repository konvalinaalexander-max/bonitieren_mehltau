"""Erzeugt die Vorlagen im Ordner vorlagen/ (ohne Betriebsdaten).

- pilot_bonitur_vorlage.xlsx: Blätter Bonitur, Wiederholung, Lichttest, Anleitung
  (Spalten genau wie in docs/SPEZIFIKATION.md, Abschnitt Pilot), mit Auswahllisten.
- projekttagebuch.xlsx: Projekttagebuch (Leitfaden 1.6) mit einer Beispielzeile.
- bonitur_bogen_pilot.pdf: Kopiervorlage Bonitur-Bogen (Seite 37 des Leitfadens).

Die Topf-Karten und die Maßstab-Karte entstehen aus der Druckseite etiketten/
(Browser: Drucken -> „Als PDF speichern“).

Aufruf:  python3 werkzeuge/vorlagen-erzeugen.py      (braucht openpyxl und pypdf)
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pypdf import PdfReader, PdfWriter

WURZEL = Path(__file__).resolve().parent.parent
ZIEL = WURZEL / 'vorlagen'
ZEILEN = 400  # vorbereitete Zeilen mit Auswahllisten
NOTEN = '0,1,2,3,4'  # Vorschlag aus dem Leitfaden; bei anderer Skala hier und in der Excel anpassen

KOPF = Font(bold=True, color='FFFFFF')
KOPF_FUELLUNG = PatternFill('solid', fgColor='2F6B3A')
HINWEIS = Font(italic=True, color='6B7570')


def blatt(wb, name, spalten, breiten, erstes=False):
    ws = wb.active if erstes else wb.create_sheet(name)
    ws.title = name
    ws.append(spalten)
    for i, b in enumerate(breiten, start=1):
        zelle = ws.cell(row=1, column=i)
        zelle.font = KOPF
        zelle.fill = KOPF_FUELLUNG
        zelle.alignment = Alignment(vertical='center')
        ws.column_dimensions[zelle.column_letter].width = b
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{ws.cell(row=1, column=len(spalten)).column_letter}{ZEILEN + 1}'
    return ws


def auswahl(ws, spalte, liste, titel, text):
    dv = DataValidation(type='list', formula1=f'"{liste}"', allow_blank=True,
                        showErrorMessage=True, errorTitle=titel, error=text)
    ws.add_data_validation(dv)
    dv.add(f'{spalte}2:{spalte}{ZEILEN + 1}')


def zahl(ws, spalte, art, von, bis, titel, text):
    dv = DataValidation(type=art, operator='between', formula1=str(von), formula2=str(bis), allow_blank=True,
                        showErrorMessage=True, errorTitle=titel, error=text)
    ws.add_data_validation(dv)
    dv.add(f'{spalte}2:{spalte}{ZEILEN + 1}')


def excel():
    wb = Workbook()
    wb.properties.creator = 'Mehltau-Bonitur (Vorlage)'
    wb.properties.lastModifiedBy = ''
    wb.properties.title = 'Pilot-Bonitur'

    # Blatt Bonitur: eine Zeile je Topf (Leitfaden Tabelle 6.2)
    ws = blatt(wb, 'Bonitur',
               ['topf_nr', 'datum', 'sorte', 'satz', 'behandlung', 'note_A', 'note_B',
                'prozent_A', 'prozent_B', 'sporen_unten', 'foto_datei', 'bemerkung'],
               [10, 12, 14, 8, 14, 9, 9, 11, 11, 13, 28, 36], erstes=True)
    for i in range(1, 201):
        ws.cell(row=i + 1, column=1, value=f'P{i:03d}')
    for r in range(2, ZEILEN + 2):
        ws.cell(row=r, column=2).number_format = 'DD.MM.YYYY'
    auswahl(ws, 'F', NOTEN, 'Note', 'Bitte eine ganze Note der Skala wählen (0 bis 4).')
    auswahl(ws, 'G', NOTEN, 'Note', 'Bitte eine ganze Note der Skala wählen (0 bis 4).')
    zahl(ws, 'H', 'decimal', 0, 100, 'Befall %', 'Befall in Prozent: eine Zahl von 0 bis 100.')
    zahl(ws, 'I', 'decimal', 0, 100, 'Befall %', 'Befall in Prozent: eine Zahl von 0 bis 100.')
    auswahl(ws, 'J', 'ja,nein', 'Sporen unten', 'Bitte „ja“ oder „nein“ wählen.')

    # Blatt Wiederholung: 5 Töpfe je 5 Durchgänge
    ws = blatt(wb, 'Wiederholung', ['topf_nr', 'durchgang', 'foto_datei', 'bemerkung'], [10, 11, 28, 40])
    zahl(ws, 'B', 'whole', 1, 5, 'Durchgang', 'Durchgang 1 bis 5.')

    # Blatt Lichttest: 1 Topf, morgens/mittags/abends, Hallenlicht an/aus
    ws = blatt(wb, 'Lichttest', ['topf_nr', 'zeitpunkt', 'hallenlicht', 'foto_datei', 'bemerkung'], [10, 12, 12, 28, 40])
    for i, (zeit, licht) in enumerate([(z, l) for z in ('morgens', 'mittags', 'abends') for l in ('an', 'aus')], start=2):
        ws.cell(row=i, column=2, value=zeit)
        ws.cell(row=i, column=3, value=licht)
    auswahl(ws, 'B', 'morgens,mittags,abends', 'Zeitpunkt', 'morgens, mittags oder abends')
    auswahl(ws, 'C', 'an,aus', 'Hallenlicht', 'an oder aus')

    # Blatt Anleitung
    ws = wb.create_sheet('Anleitung')
    ws.column_dimensions['A'].width = 110
    texte = [
        ('Pilot-Bonitur – so füllst du die Datei aus (Leitfaden Kapitel 6)', Font(bold=True, size=14, color='2F6B3A')),
        ('', None),
        ('Blatt „Bonitur“: eine Zeile je Topf. Die Nummern P001–P200 sind vorbereitet; Zeilen ohne Noten und Foto überspringt die Werkstatt.', None),
        ('Spaltennamen in Zeile 1 nicht ändern – die Analyse-Werkstatt liest die Blätter „Bonitur“, „Wiederholung“ und „Lichttest“ über diese Namen.', None),
        ('note_A / note_B: Note von Person A bzw. B (ganze Zahl). Die Auswahlliste enthält den Vorschlag 0–4; habt ihr in Etappe 0 eine andere Skala', None),
        ('   festgelegt, die Liste anpassen: Spalte markieren → Daten → Datenüberprüfung.', None),
        ('prozent_A / prozent_B: geschätzter Befall in % (gelbe + braune Blattfläche an der ganzen Pflanze), 0 bis 100.', None),
        ('sporen_unten: ja/nein – Sporenrasen auf der Blattunterseite (nur die eingeteilte Person, nach dem Foto).', None),
        ('foto_datei: Dateiname des Box-Fotos, z. B. BOX_20261012_101532.jpg (steht in Open Camera bzw. im Ordner DCIM/Bonitur).', None),
        ('', None),
        ('Blatt „Wiederholung“: 5 Töpfe je 5-mal fotografiert (jedes Mal herausnehmen und neu einsetzen); durchgang 1 bis 5.', None),
        ('Blatt „Lichttest“: ein Topf morgens, mittags und abends, jeweils mit Hallenlicht an und aus (6 Fotos) – Zeilen sind vorbereitet.', None),
        ('', None),
        ('Keine Namen eintragen, nur Kürzel. Diese Datei bleibt auf dem Firmen-PC (Bonitur/Pilot/) und kommt nicht ins Repository.', HINWEIS),
    ]
    for i, (text, font) in enumerate(texte, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if font:
            c.font = font
    wb.active = 0
    pfad = ZIEL / 'pilot_bonitur_vorlage.xlsx'
    wb.save(pfad)
    return pfad


def tagebuch():
    """Projekttagebuch (Leitfaden 1.6): eine Zeile je Arbeitsschritt."""
    wb = Workbook()
    wb.properties.creator = 'Mehltau-Bonitur (Vorlage)'
    wb.properties.lastModifiedBy = ''
    wb.properties.title = 'Projekttagebuch'
    ws = blatt(wb, 'Tagebuch', ['datum', 'etappe', 'erledigt', 'entscheidung', 'offene_fragen', 'naechster_schritt'],
               [12, 8, 50, 40, 45, 40], erstes=True)
    ws.append(['05.10.2026', '0', 'Claude hat Analyse-Kern, Werkstatt, Handy-App, Druckseite, Vorlagen und '
               'Druckteile vorab gebaut (PDF „Stand und Fragen“)', '', 'Fragen F1–F15 aus dem PDF',
               'Werkstatt-Demo ansehen, GitHub einrichten, Fragen beantworten, messen'])
    for r in range(2, ZEILEN + 2):
        for c in range(1, 7):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical='top')
    pfad = ZIEL / 'projekttagebuch.xlsx'
    wb.save(pfad)
    return pfad


def bogen():
    leitfaden = PdfReader(WURZEL / 'docs' / 'Leitfaden_Mehltau-Bonitur.pdf')
    for nr, seite in enumerate(leitfaden.pages):
        if 'Bonitur-Bogen – Pilot' in (seite.extract_text() or ''):
            w = PdfWriter()
            w.add_page(seite)
            w.add_metadata({'/Title': 'Bonitur-Bogen – Pilot', '/Author': 'Mehltau-Bonitur'})
            pfad = ZIEL / 'bonitur_bogen_pilot.pdf'
            with open(pfad, 'wb') as f:
                w.write(f)
            return pfad, nr + 1
    raise SystemExit('Seite mit dem Bonitur-Bogen nicht gefunden.')


if __name__ == '__main__':
    ZIEL.mkdir(exist_ok=True)
    print('geschrieben:', excel())
    print('geschrieben:', tagebuch())
    pfad, nr = bogen()
    print(f'geschrieben: {pfad} (Leitfaden Seite {nr})')
